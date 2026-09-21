import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from app.utils.file_utils import sanitize_filename


_MARKDOWN_HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+")
_MARKDOWN_LIST_RE = re.compile(r"^\s{0,3}[-*]\s+")
_FIELD_LABELS = {
    "正文标题",
    "正文内容",
    "标题",
}
_REFERENCE_SECTION_PREFIXES = (
    "本次参考资料",
    "参考资料",
    "内容参考",
    "写法参考",
    "检索结果",
    "资料库命中",
    "RAG",
)
_SOURCE_MARKERS = (
    "来源：",
    "来源:",
    "文件名：",
    "文件名:",
    "reference_filename",
    ".docx",
)
_LEVEL_ONE_HEADING_RE = re.compile(r"^[一二三四五六七八九十]+、")


def clean_word_export_text(text: str) -> str:
    """Remove common Markdown markers before writing plain Word paragraphs."""
    cleaned_lines: list[str] = []
    skipping_reference_block = False
    for raw_line in str(text or "").splitlines():
        line = raw_line.strip()
        if not line:
            skipping_reference_block = False
            cleaned_lines.append("")
            continue
        if line.startswith("```"):
            continue
        line = _MARKDOWN_HEADING_RE.sub("", line)
        line = _MARKDOWN_LIST_RE.sub("", line)
        line = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
        line = line.replace("**", "")
        line = line.strip()
        if not line:
            continue
        if _is_field_label(line):
            continue
        if _is_reference_section_line(line):
            skipping_reference_block = True
            continue
        if skipping_reference_block or _looks_like_source_line(line):
            continue
        cleaned_lines.append(line.strip())
    return "\n".join(cleaned_lines).strip()


def export_article_to_docx(result: dict, cover_image_path: str, structure_image_path: str, output_path: str) -> str:
    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FileExistsError(f"Export path already exists: {output}")

    document = Document()
    _configure_document(document)
    if result.get("content_format") == "full-article/v1":
        _add_full_article(document, result.get("body", ""))
    else:
        _add_title(document, result.get("recommended_title", "未命名标题"))
        _add_body_paragraphs(document, result.get("body", ""))

    document.save(output)
    return str(output)


def _add_full_article(document: Document, content: str) -> None:
    """Project canonical lines once; retain hierarchy, indentation and order."""
    from app.models.full_article import article_blocks, FullArticleContractError, text_sha
    blocks = article_blocks(content)
    if not blocks or blocks[0].level != 1:
        raise FullArticleContractError("FULL_ARTICLE_TITLE_REQUIRED")
    document.core_properties.subject = "SUCCESS_WITH_REVIEW / 待人工复核稿，非正式合格稿"
    document.core_properties.keywords = f"content_sha256={text_sha(content)};review_required=true"
    for block in blocks:
        style = "Title" if block == blocks[0] else f"Heading {max(1, block.level - 1)}" if block.level else "Normal"
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.line_spacing = Pt(28)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        if style == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(content[block.text_start:block.end])
        run.font.size = Pt(22 if style == "Title" else 16)
        run.bold = bool(block.level)
        _set_run_font(run, "方正小标宋简体" if style == "Title" else "黑体" if block.level == 2 else "楷体_GB2312" if block.level else "仿宋_GB2312")


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)


def _add_title(document: Document, raw_title: str) -> None:
    title_lines = _split_title_lines(clean_word_export_text(raw_title) or "未命名标题")
    for index, title_text in enumerate(title_lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6 if index + 1 < len(title_lines) else 14)
        run = paragraph.add_run(title_text)
        run.bold = True
        run.font.size = Pt(22 if index == 0 else 16)
        _set_run_font(run, "方正小标宋简体" if index == 0 else "仿宋_GB2312")


def _split_title_lines(title: str) -> list[str]:
    title = re.sub(r"\s+", " ", str(title or "").strip())
    if not title:
        return ["未命名标题"]
    if "——" in title:
        main, subtitle = title.split("——", 1)
        main = main.strip()
        subtitle = subtitle.strip()
        if main and subtitle:
            return [main, f"——{subtitle}"]
    match = re.search(r"\s+(在.+上的讲话)$", title)
    if match:
        main = title[: match.start()].strip()
        subtitle = match.group(1).strip()
        if main and subtitle:
            return [main, f"——{subtitle}"]
    return [title]


def _add_image_if_present(document: Document, image_path: str) -> None:
    if not image_path:
        return
    path = Path(image_path)
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.2))


def _add_body_paragraphs(document: Document, body: str) -> None:
    for raw_line in clean_word_export_text(body).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = Pt(28)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.size = Pt(16)
        _set_run_font(run, "仿宋_GB2312")
        if _is_subheading(line):
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_before = Pt(8)
            run.bold = True
            _set_run_font(run, "黑体")
        else:
            paragraph.paragraph_format.first_line_indent = Pt(32)


def _is_subheading(text: str) -> bool:
    return bool(_LEVEL_ONE_HEADING_RE.match(text.strip()))


def _is_field_label(text: str) -> bool:
    normalized = text.strip().strip("：:")
    return normalized in _FIELD_LABELS


def _is_reference_section_line(text: str) -> bool:
    normalized = text.strip().strip("：:")
    return any(normalized.startswith(prefix) for prefix in _REFERENCE_SECTION_PREFIXES)


def _looks_like_source_line(text: str) -> bool:
    return any(marker in text for marker in _SOURCE_MARKERS)


def _set_run_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
