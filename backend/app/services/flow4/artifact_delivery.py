"""Deterministic TXT/DOCX delivery for validated FLOW4 render results."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.models.flow4 import Flow4RenderResult, canonical_payload_sha256
from app.models.flow4_artifact import (
    ArtifactDeliveryResult,
    ArtifactEvidenceTraceItem,
    ArtifactFileType,
    ArtifactManifest,
    ArtifactManifestEntry,
    ArtifactMetadata,
    ArtifactParagraph,
)
from app.utils.file_utils import sanitize_title_filename


_MANIFEST_VERSION = "flow4-artifact-manifest/v2"
_PARAGRAPH_MANIFEST_VERSION = "flow4-artifact-manifest/v3"
_PACKAGE_ID_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}")
_FLOW4_DOCX_AUTHOR = "FLOW4 Artifact Delivery Layer"


class ArtifactDeliveryError(ValueError):
    """Raised when a render result cannot be safely delivered."""


def deliver_full_article_artifacts(render_result, output_dir):
    """Deliver a fresh version-bound package without overwriting existing files."""
    from app.models.full_article import canonical_json, text_sha, article_blocks
    from app.services.flow4.renderer import render_full_article
    from app.services.export.word_exporter import export_article_to_docx
    snapshot = render_full_article(render_result.handoff)
    if snapshot != render_result:
        raise ArtifactDeliveryError("full article render projection changed")
    root = Path(output_dir).resolve() / f"FULL-{snapshot.handoff.candidate.digest[:16]}-{uuid.uuid4().hex[:12]}"
    root.mkdir(parents=True, exist_ok=False)
    txt, docx = root / "article.txt", root / "article.docx"
    _write_exact_txt(txt, snapshot.final_text)
    export_article_to_docx({"body": snapshot.final_text, "content_format": "full-article/v1"}, "", "", str(docx))
    document = Document(docx)
    blocks = article_blocks(snapshot.final_text)
    expected_text = [snapshot.final_text[b.text_start:b.end] for b in blocks]
    expected_styles = ["Title" if b == blocks[0] else f"Heading {max(1, b.level - 1)}" if b.level else "Normal" for b in blocks]
    if ([p.text for p in document.paragraphs] != expected_text
            or [p.style.name for p in document.paragraphs] != expected_styles
            or txt.read_bytes() != snapshot.final_text.encode("utf-8")):
        raise ArtifactDeliveryError("full article artifact round trip failed")
    manifest = {"contract": "flow4-full-article-artifact/v1",
        "candidate_digest": snapshot.handoff.candidate.digest,
        "content_sha256": text_sha(snapshot.final_text),
        "evidence_sha256": snapshot.handoff.candidate.evidence_sha256,
        "fact_report_sha256": snapshot.handoff.fact_report.result_sha256,
        "quality_report_sha256": snapshot.handoff.quality_report.result_sha256,
        "review_status": "SUCCESS_WITH_REVIEW", "projection": snapshot.projection.model_dump(),
        "artifacts": [{"path": str(p), "sha256": _sha256_file(p), "size": p.stat().st_size} for p in (txt, docx)],
        "docx_text_order_styles_verified": True}
    (root / "artifact_manifest.json").write_text(canonical_json(manifest), encoding="utf-8")
    return manifest


def calculate_render_result_sha256(render_result: Flow4RenderResult) -> str:
    """Return the canonical SHA-256 for the complete renderer output."""

    snapshot = _validated_render_snapshot(render_result)
    return canonical_payload_sha256(snapshot.model_dump(mode="json"))


def deliver_flow4_artifacts(
    render_result: Flow4RenderResult,
    output_dir: str | Path,
    *,
    formats: Iterable[ArtifactFileType] = ("txt", "docx"),
    output_name: str | None = None,
    package_id: str | None = None,
    generated_time: datetime | None = None,
) -> ArtifactDeliveryResult:
    """Create a traceable, integrity-checked FLOW4 artifact package."""

    snapshot = _validated_render_snapshot(render_result)
    selected_formats = _normalize_formats(formats)
    source_render_hash = calculate_render_result_sha256(snapshot)
    created_at = generated_time or datetime.now(timezone.utc)
    if created_at.tzinfo is None or created_at.utcoffset() is None:
        raise ArtifactDeliveryError("generated_time must include a timezone")

    resolved_package_id = package_id or (
        f"FLOW4-{source_render_hash[:16]}-{uuid.uuid4().hex[:12]}"
    )
    if not _PACKAGE_ID_RE.fullmatch(resolved_package_id):
        raise ArtifactDeliveryError(
            "package_id must contain only letters, digits, dot, underscore, or hyphen"
        )

    root = Path(output_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    package_path = root / resolved_package_id
    if package_path.exists():
        raise FileExistsError(f"artifact package already exists: {package_path}")

    safe_name = sanitize_title_filename(
        output_name or snapshot.title,
        fallback="flow4_artifact",
    )
    temp_package = Path(
        tempfile.mkdtemp(prefix=f".{resolved_package_id}.", dir=root)
    )
    committed = False

    try:
        entries: list[ArtifactManifestEntry] = []
        for file_type in selected_formats:
            artifact_id = f"{resolved_package_id}-{file_type.upper()}"
            file_name = f"{safe_name}.{file_type}"
            temp_file_path = temp_package / file_name
            final_file_path = package_path / file_name

            if file_type == "txt":
                _write_exact_txt(temp_file_path, snapshot.final_text)
                if temp_file_path.read_bytes() != snapshot.final_text.encode("utf-8"):
                    raise ArtifactDeliveryError("TXT byte content changed during delivery")
            else:
                _write_docx(temp_file_path, snapshot, artifact_id, source_render_hash, created_at)
                _verify_docx_structure(temp_file_path, snapshot)

            file_size = temp_file_path.stat().st_size
            file_sha256 = _sha256_file(temp_file_path)
            evidence_trace = _build_evidence_trace(snapshot)
            metadata = ArtifactMetadata(
                artifact_id=artifact_id,
                source_render_hash=source_render_hash,
                generated_time=created_at,
                file_type=file_type,
                file_path=str(final_file_path),
                file_size=file_size,
                file_sha256=file_sha256,
                validation_summary=snapshot.validation_summary,
                review_status=snapshot.validation_summary.risk_status,
                render_trace=snapshot.render_trace,
                evidence_trace=evidence_trace,
                paragraphs=tuple(
                    ArtifactParagraph(
                        **paragraph.model_dump(), section_id=section.section_id,
                        sentence_texts=tuple(section.sentences[index].text for index in paragraph.sentence_indexes),
                    )
                    for section in snapshot.sections for paragraph in section.paragraphs
                ),
            )
            metadata_name = f"{file_name}.metadata.json"
            temp_metadata_path = temp_package / metadata_name
            final_metadata_path = package_path / metadata_name
            _write_json(temp_metadata_path, metadata.model_dump(mode="json"))

            entries.append(
                ArtifactManifestEntry(
                    artifact_id=artifact_id,
                    file_type=file_type,
                    file_path=str(final_file_path),
                    file_size=file_size,
                    sha256=file_sha256,
                    source_render_hash=source_render_hash,
                    metadata_path=str(final_metadata_path),
                    metadata_size=temp_metadata_path.stat().st_size,
                    metadata_sha256=_sha256_file(temp_metadata_path),
                )
            )

        manifest = ArtifactManifest(
            manifest_version=(_PARAGRAPH_MANIFEST_VERSION if any(section.paragraphs for section in snapshot.sections) else _MANIFEST_VERSION),
            package_id=resolved_package_id,
            source_render_hash=source_render_hash,
            generated_time=created_at,
            review_status=snapshot.validation_summary.risk_status,
            validation_summary=snapshot.validation_summary,
            artifacts=tuple(entries),
        )
        temp_manifest_path = temp_package / "artifact_manifest.json"
        _write_json(temp_manifest_path, manifest.model_dump(mode="json"))
        _verify_staged_package(temp_package, manifest)

        os.replace(temp_package, package_path)
        committed = True
        manifest_path = package_path / "artifact_manifest.json"
        verified_manifest = verify_artifact_manifest(manifest_path)
        return ArtifactDeliveryResult(
            package_id=resolved_package_id,
            package_path=str(package_path),
            manifest_path=str(manifest_path),
            source_render_hash=source_render_hash,
            artifacts=verified_manifest.artifacts,
        )
    finally:
        if not committed and temp_package.exists():
            shutil.rmtree(temp_package)


def verify_artifact_manifest(manifest_path: str | Path) -> ArtifactManifest:
    """Recheck every artifact and metadata sidecar against a saved manifest."""

    path = Path(manifest_path).resolve()
    manifest = ArtifactManifest.model_validate_json(path.read_text(encoding="utf-8"))
    package_path = path.parent
    if manifest.package_id != package_path.name:
        raise ArtifactDeliveryError("manifest package_id does not match its directory")

    for entry in manifest.artifacts:
        artifact_path = _resolve_manifest_member(entry.file_path, package_path)
        metadata_path = _resolve_manifest_member(entry.metadata_path, package_path)
        _verify_file_record(
            artifact_path,
            expected_size=entry.file_size,
            expected_sha256=entry.sha256,
        )
        _verify_file_record(
            metadata_path,
            expected_size=entry.metadata_size,
            expected_sha256=entry.metadata_sha256,
        )

        metadata = ArtifactMetadata.model_validate_json(
            metadata_path.read_text(encoding="utf-8")
        )
        if manifest.manifest_version == _MANIFEST_VERSION and metadata.paragraphs:
            raise ArtifactDeliveryError("paragraph snapshots require artifact manifest v3")
        if manifest.manifest_version == _PARAGRAPH_MANIFEST_VERSION and not metadata.paragraphs:
            raise ArtifactDeliveryError("PARAGRAPH_ARTIFACT_SNAPSHOT_MISSING")
        if (
            metadata.artifact_id != entry.artifact_id
            or metadata.file_type != entry.file_type
            or Path(metadata.file_path).resolve() != artifact_path
            or metadata.file_size != entry.file_size
            or metadata.file_sha256 != entry.sha256
            or metadata.source_render_hash != manifest.source_render_hash
            or entry.source_render_hash != manifest.source_render_hash
            or metadata.review_status != manifest.review_status
            or metadata.validation_summary != manifest.validation_summary
            or metadata.generated_time != manifest.generated_time
        ):
            raise ArtifactDeliveryError(
                f"metadata does not match manifest entry: {entry.artifact_id}"
            )
    return manifest


def read_artifact_metadata(
    manifest_path: str | Path, *, artifact_id: str,
) -> ArtifactMetadata:
    """Read verified sentence/paragraph ownership from its persisted artifact.

    Legacy v2 metadata remains readable and has no declared paragraph snapshot.
    A bare TXT/DOCX cannot recover IDs; its verified metadata is required.
    """
    manifest = verify_artifact_manifest(manifest_path)
    matches = [entry for entry in manifest.artifacts if entry.artifact_id == artifact_id]
    if len(matches) != 1:
        raise ArtifactDeliveryError("artifact_id is absent or ambiguous")
    return ArtifactMetadata.model_validate_json(
        Path(matches[0].metadata_path).read_text(encoding="utf-8")
    )


def _validated_render_snapshot(render_result: Flow4RenderResult) -> Flow4RenderResult:
    if not isinstance(render_result, Flow4RenderResult):
        raise TypeError("artifact delivery input must be a Flow4RenderResult")
    try:
        snapshot = Flow4RenderResult.model_validate(render_result.model_dump(mode="python"))
    except Exception as exc:
        raise ArtifactDeliveryError(f"invalid Flow4RenderResult: {exc}") from exc
    if not snapshot.render_trace:
        raise ArtifactDeliveryError("Flow4RenderResult has no rendered sentence trace")
    return snapshot


def _normalize_formats(formats: Iterable[ArtifactFileType]) -> tuple[ArtifactFileType, ...]:
    if isinstance(formats, str):
        candidates = (formats,)
    else:
        candidates = tuple(formats)
    normalized: list[ArtifactFileType] = []
    for candidate in candidates:
        if candidate not in {"txt", "docx"}:
            raise ArtifactDeliveryError(f"unsupported artifact format: {candidate!r}")
        if candidate not in normalized:
            normalized.append(candidate)
    if not normalized:
        raise ArtifactDeliveryError("at least one artifact format is required")
    return tuple(normalized)


def _build_evidence_trace(
    render_result: Flow4RenderResult,
) -> tuple[ArtifactEvidenceTraceItem, ...]:
    return tuple(
        ArtifactEvidenceTraceItem(
            section_id=item.section_id,
            sentence_index=item.sentence_index,
            sentence_type=item.sentence_type,
            evidence_ids=item.evidence_ids,
            text_sha256=item.rendered_text_sha256,
            anchor_roles=item.anchor_roles,
            trace_status=item.trace_status,
            subsection_id=item.subsection_id,
            plan_step=item.plan_step,
            writing_focus_role=item.writing_focus_role,
            expansion_stage=item.expansion_stage,
            content_function=item.content_function,
            work_chain_role=item.work_chain_role,
        )
        for item in render_result.render_trace
    )


def _write_exact_txt(path: Path, text: str) -> None:
    path.write_bytes(text.encode("utf-8"))


def _write_json(path: Path, value: object) -> None:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        indent=2,
        allow_nan=False,
    )
    path.write_bytes((payload + "\n").encode("utf-8"))


def _write_docx(
    path: Path,
    render_result: Flow4RenderResult,
    artifact_id: str,
    source_render_hash: str,
    generated_time: datetime,
) -> None:
    document = Document()
    _configure_standard_business_brief(document)

    properties = document.core_properties
    properties.title = render_result.title
    properties.subject = "Deterministic FLOW4 delivery artifact"
    properties.author = _FLOW4_DOCX_AUTHOR
    properties.last_modified_by = _FLOW4_DOCX_AUTHOR
    properties.keywords = (
        f"artifact_id={artifact_id};"
        f"review_status={render_result.validation_summary.risk_status};"
        f"source_render_hash={source_render_hash}"
    )
    utc_naive = generated_time.astimezone(timezone.utc).replace(tzinfo=None)
    properties.created = utc_naive
    properties.modified = utc_naive

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(render_result.title)

    for section in render_result.sections:
        heading = document.add_paragraph(style="Heading 1")
        heading.add_run(section.heading)
        for body_line in section.body_lines:
            paragraph = document.add_paragraph(style="Normal")
            paragraph.add_run(body_line)

    document.save(path)


def _configure_standard_business_brief(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", 11, RGBColor(0x00, 0x00, 0x00))
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = document.styles["Title"]
    _set_style_font(title, "Calibri", 20, RGBColor(0x0B, 0x25, 0x45))
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.0

    heading = document.styles["Heading 1"]
    _set_style_font(heading, "Calibri", 16, RGBColor(0x2E, 0x74, 0xB5))
    heading.font.bold = True
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.line_spacing = 1.0
    heading.paragraph_format.keep_with_next = True


def _set_style_font(style, font_name: str, size: int, color: RGBColor) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)


def _verify_docx_structure(path: Path, render_result: Flow4RenderResult) -> None:
    try:
        document = Document(path)
    except Exception as exc:
        raise ArtifactDeliveryError(f"DOCX cannot be opened: {exc}") from exc

    expected_text = [render_result.title]
    expected_styles = ["Title"]
    for section in render_result.sections:
        expected_text.append(section.heading)
        expected_styles.append("Heading 1")
        for body_line in section.body_lines:
            expected_text.append(body_line)
            expected_styles.append("Normal")

    actual_text = [paragraph.text for paragraph in document.paragraphs]
    actual_styles = [paragraph.style.name for paragraph in document.paragraphs]
    if actual_text != expected_text:
        raise ArtifactDeliveryError("DOCX title, heading, or body text/order changed")
    if actual_styles != expected_styles:
        raise ArtifactDeliveryError("DOCX title or section style structure changed")


def _verify_staged_package(temp_package: Path, manifest: ArtifactManifest) -> None:
    for entry in manifest.artifacts:
        artifact_path = temp_package / Path(entry.file_path).name
        metadata_path = temp_package / Path(entry.metadata_path).name
        _verify_file_record(artifact_path, entry.file_size, entry.sha256)
        _verify_file_record(
            metadata_path,
            entry.metadata_size,
            entry.metadata_sha256,
        )
        metadata = ArtifactMetadata.model_validate_json(
            metadata_path.read_text(encoding="utf-8")
        )
        if (
            metadata.source_render_hash != manifest.source_render_hash
            or metadata.review_status != manifest.review_status
            or metadata.validation_summary != manifest.validation_summary
        ):
            raise ArtifactDeliveryError("staged metadata lost FLOW4 audit state")


def _resolve_manifest_member(value: str, package_path: Path) -> Path:
    resolved = Path(value).resolve()
    if not resolved.is_relative_to(package_path):
        raise ArtifactDeliveryError("manifest file path escapes its artifact package")
    return resolved


def _verify_file_record(
    path: Path,
    expected_size: int,
    expected_sha256: str,
) -> None:
    if not path.is_file():
        raise ArtifactDeliveryError(f"artifact file is missing: {path}")
    if path.stat().st_size != expected_size:
        raise ArtifactDeliveryError(f"artifact size mismatch: {path}")
    if _sha256_file(path) != expected_sha256:
        raise ArtifactDeliveryError(f"artifact SHA-256 mismatch: {path}")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
