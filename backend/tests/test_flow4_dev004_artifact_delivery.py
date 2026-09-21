from __future__ import annotations

import ast
import hashlib
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document

from app.models.flow4 import Flow4RenderedSection, Flow4RenderResult
from app.models.flow4_artifact import ArtifactMetadata
from app.services.flow4.artifact_delivery import (
    ArtifactDeliveryError,
    calculate_render_result_sha256,
    deliver_flow4_artifacts,
    verify_artifact_manifest,
)
from app.services.flow4.handoff import build_flow4_input, create_flow3_output_envelope
from app.services.flow4.integrity_gate import assert_flow4_boundary_isolated
from app.services.flow4.renderer import render_flow4_input


_GENERATED_TIME = datetime(2026, 7, 27, 9, 30, tzinfo=timezone.utc)


def _gated_input():
    envelope = create_flow3_output_envelope(
        parsed_structured_draft={
            "title": "大型活动安保经验材料",
            "sections": [
                {
                    "section_id": "SEC-01",
                    "heading": "一、强化风险前置研判",
                    "sentences": [
                        {
                            "sentence_type": "FACT",
                            "evidence_id": "EV-001",
                            "text": "县委专题听取安保工作汇报。",
                        },
                        {
                            "sentence_type": "ANALYSIS",
                            "evidence_id": "",
                            "text": "风险前置研判有助于增强工作主动性。",
                        },
                    ],
                },
                {
                    "section_id": "SEC-02",
                    "heading": "二、健全协同工作机制",
                    "sentences": [
                        {
                            "sentence_type": "FACT",
                            "evidence_id": "EV-002",
                            "text": "县委召开多部门联席会议。",
                        }
                    ],
                },
            ],
        },
        content_plan={
            "schema_version": "flow3-v3.3-content-plan-1",
            "plan_id": "PLAN-DEV004",
            "sections": [{"section_id": "SEC-01"}, {"section_id": "SEC-02"}],
            "target_total": 800,
        },
        fact_anchor_pack={
            "schema_version": "flow3-v3.3-fact-anchor-1",
            "fact_anchors": [
                {
                    "evidence_id": "EV-001",
                    "section_id": "SEC-01",
                    "role": "core",
                    "canonical_subject": "县委",
                    "canonical_action": "专题听取",
                    "canonical_object": "安保工作汇报",
                    "exact_values": {},
                    "forbidden_detail_types": [],
                },
                {
                    "evidence_id": "EV-002",
                    "section_id": "SEC-02",
                    "role": "supporting",
                    "canonical_subject": "县委",
                    "canonical_action": "召开",
                    "canonical_object": "多部门联席会议",
                    "exact_values": {},
                    "forbidden_detail_types": [],
                },
            ],
            "analysis_support_pack": [],
            "source_span_dedup_report": {
                "passed": True,
                "source_span_cross_section_conflict_count": 0,
            },
        },
        validation_reports={
            "structured": {
                "passed": True,
                "invalid_evidence_reference_count": 0,
            },
            "fact": {
                "passed": True,
                "unsupported_count": 0,
                "out_of_plan_count": 0,
            },
            "analysis": {
                "passed": False,
                "analysis_fact_leak_count": 2,
                "leaks": [{"sentence_index": 0}, {"sentence_index": 1}],
                "checker_modified_response": False,
            },
        },
        task={
            "title": "大型活动安保经验材料",
            "document_type": "经验交流材料",
            "target_words": 800,
            "export_format": "docx",
        },
    )
    return build_flow4_input(
        envelope,
        run_id="FLOW3-RUN-DEV004",
        flow3_status="FAIL_ANALYSIS_FACT_LEAK",
        handoff_id="HANDOFF-DEV004",
        created_at=datetime(2026, 7, 27, 8, 0, tzinfo=timezone.utc),
        render_options={"include_title_in_plain_text": True},
    )


def _render_result():
    return render_flow4_input(_gated_input())


def _deliver(tmp_path: Path, *, formats=("txt", "docx")):
    return deliver_flow4_artifacts(
        _render_result(),
        tmp_path,
        formats=formats,
        package_id="FLOW4-DEV004-ARTIFACT",
        generated_time=_GENERATED_TIME,
    )


def test_txt_content_is_byte_exact_renderer_final_text(tmp_path):
    render_result = _render_result()
    delivery = deliver_flow4_artifacts(
        render_result,
        tmp_path,
        formats=("txt",),
        package_id="FLOW4-DEV004-TXT",
        generated_time=_GENERATED_TIME,
    )
    txt_entry = delivery.artifacts[0]
    txt_path = Path(txt_entry.file_path)

    assert txt_path.read_bytes() == render_result.final_text.encode("utf-8")
    assert txt_path.read_text(encoding="utf-8") == render_result.final_text


def test_render_result_is_detached_from_later_mutation_of_gate_input(tmp_path):
    source = _gated_input()
    render_result = render_flow4_input(source)
    original_text = render_result.final_text
    source.parsed_structured_draft.sections[0].sentences[0].text = "后续篡改。"

    delivery = deliver_flow4_artifacts(
        render_result,
        tmp_path,
        formats=("txt",),
        package_id="FLOW4-DEV004-DETACHED",
        generated_time=_GENERATED_TIME,
    )

    assert Path(delivery.artifacts[0].file_path).read_text(encoding="utf-8") == original_text


def test_docx_opens_and_preserves_title_sections_and_sentence_order(tmp_path):
    render_result = _render_result()
    delivery = deliver_flow4_artifacts(
        render_result,
        tmp_path,
        formats=("docx",),
        package_id="FLOW4-DEV004-DOCX",
        generated_time=_GENERATED_TIME,
    )
    document = Document(delivery.artifacts[0].file_path)

    assert [paragraph.text for paragraph in document.paragraphs] == [
        "大型活动安保经验材料",
        "一、强化风险前置研判",
        "县委专题听取安保工作汇报。",
        "风险前置研判有助于增强工作主动性。",
        "二、健全协同工作机制",
        "县委召开多部门联席会议。",
    ]
    assert [paragraph.style.name for paragraph in document.paragraphs] == [
        "Title",
        "Heading 1",
        "Normal",
        "Normal",
        "Heading 1",
        "Normal",
    ]


def test_docx_uses_fixed_business_brief_layout_and_audit_properties(tmp_path):
    delivery = _deliver(tmp_path, formats=("docx",))
    entry = delivery.artifacts[0]
    document = Document(entry.file_path)
    section = document.sections[0]

    assert section.page_width.inches == pytest.approx(8.5, abs=0.01)
    assert section.page_height.inches == pytest.approx(11, abs=0.01)
    assert section.top_margin.inches == pytest.approx(1, abs=0.01)
    assert section.right_margin.inches == pytest.approx(1, abs=0.01)
    assert section.bottom_margin.inches == pytest.approx(1, abs=0.01)
    assert section.left_margin.inches == pytest.approx(1, abs=0.01)

    normal = document.styles["Normal"]
    heading = document.styles["Heading 1"]
    assert normal.font.name == "Calibri"
    assert normal.font.size.pt == pytest.approx(11)
    assert normal.paragraph_format.space_after.pt == pytest.approx(6)
    assert normal.paragraph_format.line_spacing == pytest.approx(1.1)
    assert heading.font.name == "Calibri"
    assert heading.font.size.pt == pytest.approx(16)
    assert heading.paragraph_format.space_before.pt == pytest.approx(16)
    assert heading.paragraph_format.space_after.pt == pytest.approx(8)

    assert "review_status=REVIEW_REQUIRED" in document.core_properties.keywords
    assert f"source_render_hash={delivery.source_render_hash}" in (
        document.core_properties.keywords
    )
    with zipfile.ZipFile(entry.file_path) as archive:
        assert archive.testzip() is None


def test_manifest_tracks_render_hash_and_verifies_file_hashes(tmp_path):
    render_result = _render_result()
    delivery = _deliver(tmp_path)
    manifest = verify_artifact_manifest(delivery.manifest_path)

    assert manifest.source_render_hash == calculate_render_result_sha256(render_result)
    assert {entry.file_type for entry in manifest.artifacts} == {"txt", "docx"}
    for entry in manifest.artifacts:
        artifact_path = Path(entry.file_path)
        assert entry.file_size == artifact_path.stat().st_size
        assert entry.sha256 == hashlib.sha256(artifact_path.read_bytes()).hexdigest()
        assert entry.source_render_hash == manifest.source_render_hash


def test_review_required_validation_and_traces_are_preserved_in_metadata(tmp_path):
    delivery = _deliver(tmp_path, formats=("txt",))
    entry = delivery.artifacts[0]
    metadata = ArtifactMetadata.model_validate_json(
        Path(entry.metadata_path).read_text(encoding="utf-8")
    )

    assert metadata.review_status == "REVIEW_REQUIRED"
    assert metadata.validation_summary.risk_status == "REVIEW_REQUIRED"
    assert metadata.validation_summary.analysis["analysis_fact_leak_count"] == 2
    assert metadata.validation_summary.analysis["passed"] is False
    assert metadata.render_trace
    assert [
        (item.sentence_type, item.evidence_ids, item.anchor_roles, item.trace_status)
        for item in metadata.evidence_trace
    ] == [
        ("FACT", ("EV-001",), ("core",), "MAPPED"),
        ("ANALYSIS", (), (), "ANALYSIS_UNANCHORED"),
        ("FACT", ("EV-002",), ("supporting",), "MAPPED"),
    ]


def test_non_renderer_and_empty_bypassed_results_are_rejected(tmp_path):
    with pytest.raises(TypeError, match="Flow4RenderResult"):
        deliver_flow4_artifacts(_gated_input(), tmp_path)  # type: ignore[arg-type]

    bypassed = Flow4RenderResult.model_construct(
        title="",
        include_title_in_final_text=True,
        sections=(),
        final_text="",
        render_trace=(),
        validation_summary=None,
    )
    with pytest.raises(ArtifactDeliveryError, match="invalid Flow4RenderResult"):
        deliver_flow4_artifacts(bypassed, tmp_path)

    render_result = _render_result()
    headings_only = render_result.model_copy(
        update={
            "sections": (
                Flow4RenderedSection(
                    section_id="SEC-EMPTY",
                    heading="一、空章节",
                    sentences=(),
                ),
            ),
            "final_text": "大型活动安保经验材料\n\n一、空章节",
            "render_trace": (),
        }
    )
    with pytest.raises(ArtifactDeliveryError, match="no rendered sentence trace"):
        deliver_flow4_artifacts(headings_only, tmp_path)


def test_manifest_verifier_rejects_tampered_delivery_file(tmp_path):
    delivery = _deliver(tmp_path, formats=("txt",))
    txt_path = Path(delivery.artifacts[0].file_path)
    txt_path.write_bytes(txt_path.read_bytes() + b"tampered")

    with pytest.raises(ArtifactDeliveryError, match="size mismatch"):
        verify_artifact_manifest(delivery.manifest_path)


def test_delivery_boundary_has_no_provider_rag_repair_or_text_cleanup_calls():
    delivery_path = (
        Path(__file__).parents[1]
        / "app"
        / "services"
        / "flow4"
        / "artifact_delivery.py"
    )
    assert_flow4_boundary_isolated([delivery_path])

    tree = ast.parse(delivery_path.read_text(encoding="utf-8"))
    imported_modules = {
        node.module or ""
        for node in ast.walk(tree)
        if isinstance(node, ast.ImportFrom)
    }
    assert all(not module.startswith("app.services.generation") for module in imported_modules)
    assert all(not module.startswith("app.services.library") for module in imported_modules)
    assert all(not module.startswith("app.services.retrieval") for module in imported_modules)
    source = delivery_path.read_text(encoding="utf-8").lower()
    for prohibited in (
        "_run_hybrid_fact_safe_gate",
        "clean_word_export_text",
        "get_llm_provider",
        "generate_text",
        "embedding",
    ):
        assert prohibited not in source


def test_manifest_json_contains_required_delivery_fields(tmp_path):
    delivery = _deliver(tmp_path, formats=("txt",))
    payload = json.loads(Path(delivery.manifest_path).read_text(encoding="utf-8"))
    entry = payload["artifacts"][0]

    assert payload["review_status"] == "REVIEW_REQUIRED"
    assert payload["source_render_hash"] == delivery.source_render_hash
    assert entry["file_path"]
    assert entry["file_size"] > 0
    assert len(entry["sha256"]) == 64
    assert entry["source_render_hash"] == delivery.source_render_hash
